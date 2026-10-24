"""Testes do Organizador Usinas — foco no bug da NF corrigido."""
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest

from contas_a_pagar_suite.modules.usinas import ProcessadorUsinas, ResultadoUsinas, exportar

CFG = {
    "padronizacao_empresa": {
        "SOLUÇÕES": ["SOLUCOES", "SOLUCOES USIMINAS", "SOLUCOES EM ACO"],
        "ARMAR":    ["ARCELORMITTAL PLANOS", "ARCELORMITTAL",
                     "ARCELOR", "ARMAR", "MITTAL"],
        "CSN":      ["COMPANHIA SIDERURGICA", "SIDERURGICA NACIONAL", "CSN"],
        "USIMINAS": ["USIMINAS", "USINA SIDERURGICA",
                     "USINAS SIDERURGICA", "USINAS"],
    },
    "nf_digitos_sufixo_padrao": 2,
    "nf_regras_especiais": [
        {"nf_termina_com": "001", "digitos_sufixo": 3},
    ],
    "nf_tamanho_minimo": 3,
}


@pytest.fixture
def proc():
    return ProcessadorUsinas(CFG)


# ── O BUG: replace('.0') em qualquer posição ──────────────────────────
def test_nf_residuo_float_so_no_final(proc):
    # "123456.0" (float do Excel) → remove só o .0 final → corta sufixo 2
    nf, _ = proc._tratar_nf("123456.0", "USIMINAS")
    assert nf == "1234"

def test_nf_ponto_no_meio_nao_corrompe(proc):
    # Regra antiga: "190.045" → replace vira "19045" → corte → "190"
    # Agora: ponto de milhar removido corretamente → "190045" → corte → "1900"
    nf, corrompia = proc._tratar_nf("190.045", "ARMAR")
    assert nf == "1900"
    assert corrompia is True  # métrica: regra antiga teria corrompido

def test_nf_vazia_vai_para_verificar(proc):
    # Regra antiga: "nan"[:-2] == "n". Agora: rejeitada.
    assert proc._tratar_nf("nan", "CSN") == (None, False)
    assert proc._tratar_nf("", "CSN") == (None, False)
    assert proc._tratar_nf("   ", "CSN") == (None, False)

def test_nf_termina_001_corta_3(proc):
    # NF terminando em "001" (parcela de 3 díg) corta 3 — vale p/ todos os grupos
    assert proc._tratar_nf("198131001", "CSN")[0] == "198131"       # CSN 19x
    assert proc._tratar_nf("994175001", "CSN")[0] == "994175"       # CSN 99x
    assert proc._tratar_nf("7860718001", "CSN")[0] == "7860718"     # CSN 78x (10 díg)
    assert proc._tratar_nf("144337001", "USIMINAS")[0] == "144337"  # USIMINAS também

def test_nf_nao_termina_001_corta_2(proc):
    # NF terminando em "01" mas NÃO "001" → parcela de 2 díg → corta 2
    assert proc._tratar_nf("781830501", "CSN")[0] == "7818305"
    assert proc._tratar_nf("98799901", "CSN")[0] == "987999"
    assert proc._tratar_nf("88123401", "CSN")[0] == "881234"

def test_nf_curta_nao_corta_abaixo_do_minimo(proc):
    # 4 dígitos - corte 2 = 2 < mínimo 3 → NÃO corta (antes cortava sempre)
    nf, _ = proc._tratar_nf("1234", "USIMINAS")
    assert nf == "1234"


# ── Padronização de empresa ───────────────────────────────────────────
def test_padronizacao_empresas(proc):
    assert proc._padronizar_empresa("USINAS SIDERURGICAS DE MG SA") == "USIMINAS"
    assert proc._padronizar_empresa("ArcelorMittal Brasil") == "ARMAR"
    assert proc._padronizar_empresa("CIA SIDERURGICA NACIONAL") == "CSN"
    assert proc._padronizar_empresa("OUTRA EMPRESA LTDA") == "OUTRA EMPRESA LTDA"


def test_padronizacao_regra_do_alan(proc):
    """Mapeamento exato que o Alan definiu (nomes reais das planilhas)."""
    # USIMINAS — os dois nomes reais do arquivo Pasta1.xlsx
    assert proc._padronizar_empresa("USINA SIDERURGICA MG USIMINAS") == "USIMINAS"
    assert proc._padronizar_empresa(
        "USINAS SIDERURGICAS DE MINAS GERAIS") == "USIMINAS"
    # ARMAR
    for nome in ("ARCELOR", "ARCELORMITTAL", "ARCELORMITTAL PLANOS", "ARMAR"):
        assert proc._padronizar_empresa(nome) == "ARMAR"
    # CSN
    assert proc._padronizar_empresa("COMPANHIA SIDERURGICA") == "CSN"
    assert proc._padronizar_empresa(
        "COMPANHIA SIDERURGICA NACIONAL") == "CSN"


def test_padronizacao_ignora_acento(proc):
    assert proc._padronizar_empresa("USINA SIDERÚRGICA") == "USIMINAS"
    assert proc._padronizar_empresa("COMPANHIA SIDERÚRGICA") == "CSN"


def test_solucoes_tem_precedencia_sobre_usiminas(proc):
    # 'SOLUÇÕES USIMINAS' contém 'USIMINAS', mas deve virar SOLUÇÕES.
    assert proc._padronizar_empresa("SOLUCOES USIMINAS") == "SOLUÇÕES"
    assert proc._padronizar_empresa("SOLUÇÕES USIMINAS") == "SOLUÇÕES"


def test_grupos_distintos_separam(proc):
    """Um arquivo com as 3 usinas deve gerar 3 grupos distintos."""
    assert proc._padronizar_empresa("USINA SIDERURGICA MG USIMINAS") == "USIMINAS"
    assert proc._padronizar_empresa("ARCELORMITTAL PLANOS SA") == "ARMAR"
    assert proc._padronizar_empresa(
        "COMPANHIA SIDERURGICA NACIONAL CSN") == "CSN"


# ── Pipeline completo com Excel de verdade ────────────────────────────
def test_pipeline_completo(tmp_path, proc):
    import pandas as pd
    entrada = tmp_path / "usinas.xlsx"
    pd.DataFrame({
        "EMPRESA":    ["Usiminas SA", "ARCELORMITTAL", "CSN", "CSN"],
        "NF":         ["123456.0", "190.045", "195002001", ""],
        "VENCIMENTO": ["20/07/2026", "2026-07-21", "22/07/2026", "23/07/2026"],
        "VALOR":      ["1.234,56", "1000", "78.115,20", "50"],
    }).to_excel(entrada, index=False)

    r = proc.processar(entrada)
    assert r.total_original == 4
    assert r.total_ok == 3
    assert r.total_verificar == 1                       # NF vazia
    assert "NF vazia" in r.linhas_verificar[0]["MOTIVO"]
    assert r.correcoes_nf == 1                          # "190.045"
    assert r.valor_total == pytest.approx(1234.56 + 1000 + 78115.20)

    # data ISO não pode virar 6 de janeiro
    armar = [l for l in r.linhas_ok if l["EMPRESA"] == "ARMAR"][0]
    assert armar["VENCIMENTO"] == datetime(2026, 7, 21)

    saida = tmp_path / "tratado.xlsx"
    exportar(r, saida)
    tratado = pd.read_excel(saida, sheet_name="Tratado", dtype={"NF": str})
    verificar = pd.read_excel(saida, sheet_name="Verificar")
    assert len(tratado) == 3
    assert len(verificar) == 1
    assert set(tratado["NF"]) == {"1234", "1900", "195002"}


def test_colunas_ausentes_erro_claro(tmp_path, proc):
    # Planilha COM cabeçalho mas sem as colunas obrigatórias → erro claro.
    import pandas as pd
    entrada = tmp_path / "ruim.xlsx"
    pd.DataFrame({"EMPRESA": ["x"], "VALOR": [1]}).to_excel(entrada, index=False)
    with pytest.raises(ValueError, match="Colunas obrigatórias"):
        proc.processar(entrada)


def test_poucas_colunas_sem_cabecalho_erro_claro(tmp_path, proc):
    # Planilha SEM cabeçalho e estreita demais para o mapa posicional do DDA.
    import pandas as pd
    entrada = tmp_path / "estreita.xlsx"
    pd.DataFrame([[1, 2], [3, 4]]).to_excel(entrada, index=False, header=False)
    with pytest.raises(ValueError, match="não consegui localizar"):
        proc.processar(entrada)


# ── Layout REAL das usinas: SEM cabeçalho e SEM filial, colado em A1 ───
# 11 colunas (0..10), medidas no arquivo real plp.xlsx:
#   [0] comprador · [1] banco · [2] CNPJ · [3] usina · [4] NF
#   [5] tipo · [6] "A pagar" · [7] VENCIMENTO · [8] data2 · [9] VALOR · [10] dup
def _lin_usina(usina, nf, venc, val):
    return ["ALFAMETAL S/A", "341", "00.000/0001", usina, nf,
            "Cobrança", "A pagar", venc, "2036-01-01", val, val]


def test_headerless_layout_usinas(tmp_path, proc):
    import pandas as pd
    dados = [
        _lin_usina("USINA SIDERURGICA MG USIMINAS", "123456.0", "20/07/2026", "1.000,00"),
        _lin_usina("ARCELORMITTAL PLANOS",          "190.045",  "21/07/2026", "2.000,00"),
        _lin_usina("COMPANHIA SIDERURGICA NACIONAL", "195002317", "22/07/2026", "3.000,00"),
    ]
    entrada = tmp_path / "colado_a1.xlsx"
    pd.DataFrame(dados).to_excel(entrada, index=False, header=False)

    r = proc.processar(entrada)
    assert r.usou_mapa_posicional is True
    assert r.total_original == 3
    assert r.total_ok == 3
    assert r.total_verificar == 0
    empresas = {l["EMPRESA"] for l in r.linhas_ok}
    assert empresas == {"USIMINAS", "ARMAR", "CSN"}


def test_usina_detectada_por_conteudo_nao_por_posicao(tmp_path, proc):
    """
    Se a usina vier numa coluna diferente da mapeada, o sistema deve
    DETECTAR pelo conteúdo e classificar cada uma no seu grupo — em vez de
    jogar tudo num grupo só (o bug do 'tudo USIMINAS').
    """
    import pandas as pd
    # usina na coluna [0] (a NF vem logo depois, em [1]); a coluna [3] do
    # mapa posicional tem lixo — a detecção por conteúdo acha a usina em [0]
    def lin(usina, nf, venc, val):
        row = [""] * 11
        row[0] = usina
        row[1] = nf
        row[3] = "LIXO NA COLUNA MAPEADA"
        row[7] = venc
        row[9] = val
        return row
    dados = [
        lin("USINA SIDERURGICA MG USIMINAS", "111111", "20/07/2026", "1.000,00"),
        lin("ARCELORMITTAL PLANOS",          "222222", "21/07/2026", "2.000,00"),
        lin("COMPANHIA SIDERURGICA NACIONAL", "333333", "22/07/2026", "3.000,00"),
    ]
    entrada = tmp_path / "layout_deslocado.xlsx"
    pd.DataFrame(dados).to_excel(entrada, index=False, header=False)

    r = proc.processar(entrada)
    assert r.usou_mapa_posicional is True
    assert r.total_ok == 3
    empresas = {l["EMPRESA"] for l in r.linhas_ok}
    assert empresas == {"USIMINAS", "ARMAR", "CSN"}   # não colapsou em 1 grupo


def test_export_nf_como_numero(tmp_path, proc):
    """A coluna NF do Excel sai como NÚMERO (sem o alerta de texto)."""
    import pandas as pd
    from openpyxl import load_workbook
    dados = [
        _lin_usina("USINA SIDERURGICA MG USIMINAS", "91203", "20/07/2026", "1.000,00"),
        _lin_usina("COMPANHIA SIDERURGICA NACIONAL", "198131001", "20/07/2026", "500,00"),
    ]
    entrada = tmp_path / "in.xlsx"
    pd.DataFrame(dados).to_excel(entrada, index=False, header=False)
    r = proc.processar(entrada)
    saida = tmp_path / "out.xlsx"
    exportar(r, saida)

    ws = load_workbook(saida)["Tratado"]
    for i in range(2, ws.max_row + 1):
        cel = ws.cell(row=i, column=2)      # coluna B = NF
        assert isinstance(cel.value, int), f"NF deveria ser número: {cel.value!r}"
        assert cel.number_format == "0"
    # a CSN 198131001 virou 198131 (corte 3)
    nfs = {ws.cell(row=i, column=2).value for i in range(2, ws.max_row + 1)}
    assert 198131 in nfs


def test_totais_por_grupo(tmp_path, proc):
    import pandas as pd
    dados = [
        _lin_usina("USINA SIDERURGICA MG USIMINAS", "111111", "20/07/2026", "1.000,00"),
        _lin_usina("USINAS SIDERURGICAS DE MG",     "222222", "20/07/2026", "500,00"),
        _lin_usina("ARCELORMITTAL PLANOS",          "333333", "21/07/2026", "2.000,00"),
        _lin_usina("COMPANHIA SIDERURGICA NACIONAL", "444444", "22/07/2026", "3.000,00"),
    ]
    entrada = tmp_path / "grupos.xlsx"
    pd.DataFrame(dados).to_excel(entrada, index=False, header=False)

    r = proc.processar(entrada)
    assert r.totais_por_grupo["USIMINAS"] == pytest.approx(1500.0)
    assert r.totais_por_grupo["ARMAR"] == pytest.approx(2000.0)
    assert r.totais_por_grupo["CSN"] == pytest.approx(3000.0)
    assert sum(r.totais_por_grupo.values()) == pytest.approx(r.valor_total)


def test_nf_milhar_texto_nao_e_corrompida(proc):
    # "198.000" (texto com ponto de MILHAR) → NF 198000 → corte 2 → "1980"
    # (antes a regex .0+$ comia ".000" e sobrava "198")
    assert proc._tratar_nf("198.000", "USIMINAS")[0] == "1980"
    # resíduo real de float continua tratado: "123456.0" → corta 2 → "1234"
    assert proc._tratar_nf("123456.0", "USIMINAS")[0] == "1234"


def test_linha_em_branco_no_meio_e_pulada(tmp_path, proc):
    """Linha vazia (NaN) no meio não vira falso problema em 'Verificar'."""
    import pandas as pd
    dados = [
        _lin_usina("USINA SIDERURGICA MG USIMINAS", "111111", "20/07/2026",
                   "1.000,00"),
        [None] * 11,                                    # linha em branco
        _lin_usina("COMPANHIA SIDERURGICA NACIONAL", "222222", "21/07/2026",
                   "2.000,00"),
    ]
    entrada = tmp_path / "com_branco.xlsx"
    pd.DataFrame(dados).to_excel(entrada, index=False, header=False)
    r = proc.processar(entrada)
    assert r.total_original == 2                        # a vazia não conta
    assert r.total_verificar == 0                       # e não vai p/ Verificar


def test_cabecalho_data_emissao_nao_rouba_vencimento(tmp_path, proc):
    """Com cabeçalho, 'DATA EMISSÃO' à esquerda não captura o VENCIMENTO."""
    import pandas as pd
    entrada = tmp_path / "hdr.xlsx"
    pd.DataFrame({
        "FORNECEDOR": ["USIMINAS SA"],
        "NF": ["123456"],
        "DATA EMISSAO": ["01/01/2026"],
        "VENCIMENTO": ["20/07/2026"],
        "VALOR": ["1.000,00"],
    }).to_excel(entrada, index=False)
    r = proc.processar(entrada)
    assert r.total_ok == 1
    from datetime import datetime as _dt
    assert r.linhas_ok[0]["VENCIMENTO"] == _dt(2026, 7, 20)   # não 01/01
