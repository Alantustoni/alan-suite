"""
test_export.py
Regressao da exportacao:
  - notas de semanas diferentes NAO somem da aba por dia (bug da 1a semana);
  - filial MG/PE sai em negrito;
  - pagamentos do banco sem nota aparecem na aba ERROS (item 5).
"""
from __future__ import annotations

import os
import sys
import tempfile
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from contas_a_pagar_suite.vendor.organizador_dda.core.config import Config
from contas_a_pagar_suite.vendor.organizador_dda.core.processor import ProcessingResult
from contas_a_pagar_suite.vendor.organizador_dda.core.banco_parser import PagamentoBanco
from contas_a_pagar_suite.vendor.organizador_dda.core.banco_cruzamento import cruzar
from contas_a_pagar_suite.vendor.organizador_dda.excel.report_generator import ExcelReportGenerator

CNPJ_WM = "90.100.600/0001-14"


def _nota(filial, empresa, cnpj, forn, prv, venc, valor):
    return {"FILIAL": filial, "EMPRESA": empresa, "CNPJ": cnpj, "FORNECEDOR": forn,
            "PRV NF": prv, "TIPO": "Cobranca", "VENC": pd.Timestamp(venc),
            "VALOR": valor, "ERROS": ""}


def _gerar():
    df = pd.DataFrame([
        _nota("MATRIZ", "ALFAMETAL MATRIZ", CNPJ_WM, "GAMMA GASOTEC", "1001", "2026-07-06", 500.0),
        _nota("MATRIZ", "ALFAMETAL MATRIZ", CNPJ_WM, "GAMMA GASOTEC", "1002", "2026-07-20", 700.0),
        _nota("MATRIZ", "ALFAMETAL MATRIZ", CNPJ_WM, "GAMMA GASOTEC", "9999", "2026-09-04", 1033.61),
        _nota("MG",     "ALFAMETAL MG",     "90.100.700/0001-40", "FORJAX", "7777", "2026-07-06", 2766.60),
    ])
    # Um pagamento do banco que NAO casa com nenhuma nota -> "sem nota"
    orfao = PagamentoBanco(
        empresa_banco="ALFAMETAL LTDA", cnpj_empresa="", favorecido="FORNECEDOR ORFAO",
        cnpj_favorecido="11111111000191", tipo_pagamento="", referencia="",
        data=date(2026, 7, 6), valor=123.45, status="",
    )
    res_cruz = cruzar(df, [orfao])

    result = ProcessingResult()
    result.df = df
    result.total_original = len(df)
    result.total_processado = len(df)
    result.error_counts = pd.Series(dtype=int)
    result.duplicatas = pd.DataFrame()
    result.resumo_remocoes = []
    result.col_map = {"VENC": "VENCIMENTO", "VALOR": "VALOR"}
    result.col_warnings = []

    saida = Path(tempfile.gettempdir()) / "teste_dda_export.xlsx"
    ExcelReportGenerator(result, Config(), resultado_cruzamento=res_cruz).gerar(saida)
    return load_workbook(saida), saida


def _valores(ws):
    return {str(c.value).strip() for row in ws.iter_rows() for c in row if c.value is not None}


def _linha_da_nota(ws, prv):
    for row in ws.iter_rows():
        for c in row:
            if str(c.value).strip() == prv:
                return c.row
    return None


def test_todas_as_semanas_aparecem():
    wb, _ = _gerar()
    vals = _valores(wb["ALFAMETAL"])
    for prv in ("1001", "1002", "9999", "7777"):
        assert prv in vals, f"nota {prv} sumiu da aba ALFAMETAL"


def test_negrito_mg():
    wb, _ = _gerar()
    ws = wb["ALFAMETAL"]
    lin_mg = _linha_da_nota(ws, "7777")
    lin_matriz = _linha_da_nota(ws, "1001")
    assert ws.cell(row=lin_mg, column=5).font.bold, "MG deveria estar em negrito"
    assert not ws.cell(row=lin_matriz, column=5).font.bold, "MATRIZ nao deveria ser negrito"


def test_pagamento_sem_nota_no_relatorio():
    wb, _ = _gerar()
    vals = _valores(wb["ERROS"])
    assert any("SEM NOTA" in v.upper() for v in vals), "faltou secao de pagamentos sem nota"
    assert "FORNECEDOR ORFAO" in vals, "pagamento orfao nao foi listado"


def test_conciliacao_ok_no_relatorio():
    wb, _ = _gerar()
    vals = _valores(wb["ERROS"])
    assert any("CONCILIA" in v.upper() for v in vals), "faltou secao de conciliacao"


def _montar_arquivo_revisado(caminho):
    """Cria um .xlsx no layout das abas ALFAMETAL/BETAVIARIA (dados em E:J)."""
    from datetime import datetime
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    header = ["EMPRESA", "CNPJ", "FORNECEDOR", "PRV NF", "VENC", "VALOR"]
    dados = {
        "ALFAMETAL": [
            ["ALFAMETAL MATRIZ", CNPJ_WM, "FORN A", "1001",
             datetime(2026, 7, 27), 500.0],
            ["ALFAMETAL MG", CNPJ_WM, "FORN B", "3230",
             datetime(2026, 7, 27), 677.5],
        ],
        "BETAVIARIA": [
            ["BETAVIARIA MATRIZ", CNPJ_WM, "FORN C", "1366",
             datetime(2026, 7, 27), 542.0],
        ],
    }
    for aba, linhas in dados.items():
        ws = wb.create_sheet(aba)
        for c, h in enumerate(header, start=5):     # cabeçalho em E:J (ignorado)
            ws.cell(row=1, column=c, value=h)
        for r, linha in enumerate(linhas, start=2):
            for c, v in enumerate(linha, start=5):
                ws.cell(row=r, column=c, value=v)
    wb.save(caminho)


def test_consolidado_negrito_mg():
    """Na aba consolidada ALFAMETAL - BTV, a filial MG sai em negrito."""
    from contas_a_pagar_suite.vendor.organizador_dda.excel.consolidador import (
        consolidar_arquivo)
    tmp = Path(tempfile.mkdtemp()) / "revisado.xlsx"
    _montar_arquivo_revisado(tmp)
    saida = consolidar_arquivo(tmp)

    # gera um arquivo NOVO (não sobrescreve o revisado) e devolve o caminho
    assert saida != tmp and saida.exists()
    assert saida.name == "revisado_consolidado.xlsx"
    assert tmp.exists()   # o revisado permanece intocado
    wb = load_workbook(saida)
    assert "ALFAMETAL - BTV" in wb.sheetnames
    ws = wb["ALFAMETAL - BTV"]
    achou_mg = achou_matriz = False
    for row in ws.iter_rows(min_col=6, max_col=6):   # coluna F = EMPRESA
        cell = row[0]
        val = str(cell.value or "").upper()
        if val == "ALFAMETAL MG":
            achou_mg = True
            assert cell.font.bold, "ALFAMETAL MG deveria estar em negrito"
        elif val == "ALFAMETAL MATRIZ":
            achou_matriz = True
            assert not cell.font.bold, "MATRIZ não deveria estar em negrito"
    assert achou_mg and achou_matriz, "linhas MG/MATRIZ não encontradas no BTV"


def _run_all():
    import traceback
    falhas = 0
    for nome, fn in sorted(globals().items()):
        if nome.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"OK  {nome}")
            except Exception:
                falhas += 1
                print(f"FALHOU  {nome}")
                traceback.print_exc()
    print("\n" + ("TODOS OS TESTES PASSARAM." if not falhas else f"{falhas} teste(s) falharam."))
    return falhas


if __name__ == "__main__":
    sys.exit(1 if _run_all() else 0)
