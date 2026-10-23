"""
widgets.py — Componentes reutilizáveis da suíte.

Tudo que aparece em mais de uma tela mora aqui: cards de métrica, badges,
tabela zebrada, barra de passos, toast, rodador de threads.
"""
from __future__ import annotations

import queue
import re
import threading
import tkinter as tk
from typing import Any, Callable, List, Optional, Sequence

import customtkinter as ctk

from ..core import theme as T


# ────────────────────────────────────────────────────────────────────── #
# Execução em thread com retorno seguro para a UI                        #
# ────────────────────────────────────────────────────────────────────── #
class ThreadRunner:
    """
    Roda trabalho pesado fora da thread da UI e entrega o resultado de
    volta via fila + after() — o padrão thread-safe do tkinter (mesma
    estratégia do DDA V7).
    """

    def __init__(self, widget: tk.Misc) -> None:
        self._widget = widget
        self._fila: "queue.Queue[tuple]" = queue.Queue()
        self._ativo = False

    def executar(
        self,
        trabalho: Callable[[], Any],
        ao_concluir: Callable[[Any], None],
        ao_falhar: Callable[[Exception], None],
    ) -> bool:
        if self._ativo:
            return False
        self._ativo = True

        def _worker() -> None:
            try:
                resultado = trabalho()
                self._fila.put(("ok", resultado))
            except Exception as exc:  # noqa: BLE001 — repassado para a UI
                self._fila.put(("erro", exc))

        threading.Thread(target=_worker, daemon=True).start()
        self._agendar_poll(ao_concluir, ao_falhar)
        return True

    def _agendar_poll(self, ao_concluir, ao_falhar) -> None:
        try:
            tipo, carga = self._fila.get_nowait()
        except queue.Empty:
            self._widget.after(80, lambda: self._agendar_poll(ao_concluir, ao_falhar))
            return
        self._ativo = False
        if tipo == "ok":
            try:
                ao_concluir(carga)
            except Exception as exc:  # noqa: BLE001
                # Erro ao MONTAR a tela não pode deixar a UI presa em
                # "processando…" — roteia para o tratador de falha da view.
                try:
                    ao_falhar(exc)
                except Exception:
                    pass
        else:
            ao_falhar(carga)

    @property
    def ocupado(self) -> bool:
        return self._ativo


# ────────────────────────────────────────────────────────────────────── #
# Cards e badges                                                         #
# ────────────────────────────────────────────────────────────────────── #
class Card(ctk.CTkFrame):
    def __init__(self, master, **kw) -> None:
        kw.setdefault("fg_color", T.CARD)
        kw.setdefault("corner_radius", T.RAIO)
        kw.setdefault("border_width", 1)
        kw.setdefault("border_color", T.BORDA)
        super().__init__(master, **kw)


class MetricCard(Card):
    """
    Card 'KPI' do design Premium: filete colorido à esquerda + rótulo
    pequeno + número grande + subtítulo (igual à prévia web).
    """

    def __init__(self, master, rotulo: str, cor_valor: str = T.INK, **kw):
        # altura compacta: o conteúdo (rótulo + número + subtítulo) ocupa
        # ~90px; sem isso o CTkFrame fica com a altura-padrão de 200px e o
        # card aparece com um vão vazio embaixo do número.
        kw.setdefault("height", 104)
        super().__init__(master, **kw)
        # cor do filete: acompanha o estado do indicador (accent por padrão)
        cor_filete = T.AZUL if cor_valor in (T.INK, T.AZUL_ESCURO) else cor_valor
        self._filete = ctk.CTkFrame(self, width=T.FILETE_W, fg_color=cor_filete,
                                    corner_radius=2)
        self._filete.pack(side="left", fill="y", padx=(6, 0), pady=8)

        corpo = ctk.CTkFrame(self, fg_color="transparent")
        corpo.pack(side="left", fill="both", expand=True)
        self._rotulo = ctk.CTkLabel(corpo, text=rotulo.upper(), font=T.F_CARD_LABEL,
                                    text_color=T.INK_2, anchor="w")
        self._rotulo.pack(fill="x", padx=12, pady=(12, 0))
        self._valor = ctk.CTkLabel(corpo, text="—", font=T.F_CARD_VALOR,
                                   text_color=cor_valor, anchor="w")
        self._valor.pack(fill="x", padx=12)
        self._sub = ctk.CTkLabel(corpo, text="", font=T.F_PEQUENA,
                                 text_color=T.INK_3, anchor="w")
        self._sub.pack(fill="x", padx=12, pady=(0, 10))

    def set(self, valor: Any, sub: str = "") -> None:
        self._valor.configure(text=str(valor))
        self._sub.configure(text=sub)


class Badge(ctk.CTkLabel):
    """Pílula de status (border-radius alto, texto bold — igual à prévia)."""

    def __init__(self, master, texto: str = "", tipo: str = "info", **kw):
        cor, bg = T.STATUS_CORES.get(tipo, T.STATUS_CORES["info"])
        super().__init__(master, text=f"  {texto}  ", font=(T.FONTE, 11, "bold"),
                         text_color=cor, fg_color=bg,
                         corner_radius=T.RAIO_PILL, height=26, **kw)

    def set(self, texto: str, tipo: str = "info") -> None:
        cor, bg = T.STATUS_CORES.get(tipo, T.STATUS_CORES["info"])
        self.configure(text=f"  {texto}  ", text_color=cor, fg_color=bg)


# ────────────────────────────────────────────────────────────────────── #
# Barra de passos (fluxo 1-2-3-4, como no DDA)                           #
# ────────────────────────────────────────────────────────────────────── #
class StepBar(ctk.CTkFrame):
    def __init__(self, master, passos: Sequence[str], **kw):
        kw.setdefault("fg_color", "transparent")
        super().__init__(master, **kw)
        self._labels: List[ctk.CTkLabel] = []
        for i, passo in enumerate(passos):
            lbl = ctk.CTkLabel(
                self, text=f"  {i + 1}. {passo}  ", font=(T.FONTE, 12, "bold"),
                corner_radius=T.RAIO_BTN, fg_color=T.CARD, text_color=T.INK_3,
                height=30,
            )
            lbl.pack(side="left", padx=(0 if i == 0 else 7, 0), ipady=2)
            self._labels.append(lbl)
        self.set_ativo(0)

    def set_ativo(self, indice: int) -> None:
        for i, lbl in enumerate(self._labels):
            if i < indice:
                lbl.configure(fg_color=T.VERDE_BG, text_color=T.VERDE)
            elif i == indice:
                lbl.configure(fg_color=T.AZUL, text_color=T.BRANCO)
            else:
                lbl.configure(fg_color=T.CARD, text_color=T.INK_3)


# ────────────────────────────────────────────────────────────────────── #
# Tabela premium — busca em tempo real + ordenar por coluna + exportar    #
# ────────────────────────────────────────────────────────────────────── #
_RE_NUM = re.compile(r"[^\d,.-]")


def _chave_ordenacao(valor: Any):
    """Ordena numérico quando dá (R$ 1.234,56 → 1234.56); senão texto."""
    s = str(valor).strip()
    limpo = _RE_NUM.sub("", s)
    if limpo and any(ch.isdigit() for ch in limpo):
        t = limpo
        if "," in t:                       # formato BR
            t = t.replace(".", "").replace(",", ".")
        try:
            return (0, float(t))
        except ValueError:
            pass
    return (1, s.lower())


class Tabela(ctk.CTkFrame):
    """
    Tabela zebrada com cabeçalho azul. Recursos opcionais (sem quebrar
    quem já usava): `busca=True` mostra um campo de filtro em tempo real,
    `exportar=True` mostra o botão "Exportar Excel", e o clique no
    cabeçalho ordena pela coluna (asc/desc). API mantida: preencher()/limpar().
    """

    def __init__(self, master, colunas: Sequence[str],
                 larguras: Optional[Sequence[int]] = None, altura: int = 220,
                 busca: bool = False, exportar: bool = False,
                 nome_export: str = "tabela",
                 alinhamentos: Optional[Sequence[str]] = None, **kw):
        kw.setdefault("fg_color", "transparent")
        super().__init__(master, **kw)
        self._colunas = list(colunas)
        self._larguras = list(larguras or [140] * len(colunas))
        # alinhamento por coluna: "w" (esquerda, padrão) ou "e" (direita p/ nº)
        self._alinhamentos = list(alinhamentos or ["w"] * len(colunas))
        if len(self._alinhamentos) < len(self._colunas):
            self._alinhamentos += ["w"] * (
                len(self._colunas) - len(self._alinhamentos))
        self._nome_export = nome_export
        self._dados: List[List[Any]] = []      # dataset completo
        self._ordem_col: Optional[int] = None
        self._ordem_desc = False
        self._linhas: List[List[ctk.CTkLabel]] = []
        self._headers: List[ctk.CTkLabel] = []

        # ── toolbar (busca + contador + exportar) ────────────────────
        if busca or exportar:
            bar = ctk.CTkFrame(self, fg_color="transparent")
            bar.pack(fill="x", pady=(0, 6))
            if busca:
                self._entry = ctk.CTkEntry(
                    bar, placeholder_text="🔎  buscar…", height=30,
                    width=240, font=T.F_PEQUENA, border_color=T.BORDA,
                    fg_color=T.CARD)
                self._entry.pack(side="left")
                self._entry.bind("<KeyRelease>", lambda _e: self._render())
            self._lbl_cont = ctk.CTkLabel(bar, text="", font=T.F_PEQUENA,
                                          text_color=T.INK_3)
            self._lbl_cont.pack(side="left", padx=10)
            if exportar:
                botao(bar, "⬇  Exportar Excel", self._exportar_excel,
                      primario=False, height=30).pack(side="right")
        else:
            self._entry = None
            self._lbl_cont = None

        # ── grade rolável ────────────────────────────────────────────
        self._grid = ctk.CTkScrollableFrame(self, height=altura,
                                            fg_color=T.CARD,
                                            corner_radius=T.RAIO)
        self._grid.pack(fill="both", expand=True)
        for c, (nome, w) in enumerate(zip(self._colunas, self._larguras)):
            # weight proporcional à largura → colunas respeitam as proporções
            # pedidas (antes era weight=1 uniforme, o que desalinhava tudo).
            self._grid.grid_columnconfigure(c, minsize=w, weight=max(1, int(w)))
            raio = T.RAIO_BTN if c in (0, len(self._colunas) - 1) else 0
            h = ctk.CTkLabel(self._grid, text=nome.upper(),
                             font=(T.FONTE, 11, "bold"), text_color=T.BRANCO,
                             fg_color=T.AZUL_ESCURO, corner_radius=raio,
                             anchor=self._alinhamentos[c])
            h.grid(row=0, column=c, sticky="nsew", padx=(0, 1), pady=(0, 3),
                   ipadx=8, ipady=5)
            h.bind("<Button-1>", lambda _e, col=c: self._ordenar(col))
            h.configure(cursor="hand2")
            self._headers.append(h)

    # ------------------------------------------------------------------ #
    def limpar(self) -> None:
        for linha in self._linhas:
            for w in linha:
                w.destroy()
        self._linhas.clear()

    def preencher(self, linhas: Sequence[Sequence[Any]],
                  max_linhas: int = 2000) -> None:
        self._dados = [list(l) for l in linhas][:max_linhas]
        self._render()

    # ------------------------------------------------------------------ #
    def _ordenar(self, col: int) -> None:
        if self._ordem_col == col:
            self._ordem_desc = not self._ordem_desc
        else:
            self._ordem_col, self._ordem_desc = col, False
        # seta no cabeçalho
        for c, h in enumerate(self._headers):
            base = self._colunas[c].upper()
            if c == col:
                h.configure(text=f"{base}  {'▼' if self._ordem_desc else '▲'}")
            else:
                h.configure(text=base)
        self._render()

    def _filtradas(self) -> List[List[Any]]:
        linhas = self._dados
        if self._entry is not None:
            termo = self._entry.get().strip().lower()
            if termo:
                linhas = [l for l in linhas
                          if any(termo in str(v).lower() for v in l)]
        if self._ordem_col is not None:
            linhas = sorted(
                linhas,
                key=lambda l: _chave_ordenacao(
                    l[self._ordem_col] if self._ordem_col < len(l) else ""),
                reverse=self._ordem_desc)
        return linhas

    def _render(self, max_visiveis: int = 400) -> None:
        self.limpar()
        linhas = self._filtradas()
        for r, linha in enumerate(linhas[:max_visiveis], start=1):
            bg = T.ZEBRA if r % 2 == 0 else T.CARD
            widgets = []
            for c, valor in enumerate(linha):
                anc = self._alinhamentos[c] if c < len(self._alinhamentos) else "w"
                lbl = ctk.CTkLabel(self._grid, text=str(valor),
                                   font=T.F_PEQUENA, text_color=T.INK,
                                   fg_color=bg, anchor=anc)
                lbl.grid(row=r, column=c, sticky="nsew", padx=(0, 1),
                         ipadx=8, ipady=3)
                widgets.append(lbl)
            self._linhas.append(widgets)
        if self._lbl_cont is not None:
            total = len(self._dados)
            vis = len(linhas)
            txt = f"{vis} de {total}" if vis != total else f"{total} linha(s)"
            if vis > max_visiveis:
                txt = f"{max_visiveis} de {vis} exibidas"
            self._lbl_cont.configure(text=txt)

    # ------------------------------------------------------------------ #
    def _exportar_excel(self) -> None:
        from tkinter import filedialog, messagebox
        from datetime import datetime
        linhas = self._filtradas()
        if not linhas:
            messagebox.showinfo("Exportar", "Não há linhas para exportar.")
            return
        caminho = filedialog.asksaveasfilename(
            title="Exportar tabela",
            defaultextension=".xlsx",
            initialfile=f"{self._nome_export}_"
                        f"{datetime.now():%Y%m%d_%H%M}.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not caminho:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = self._nome_export[:31]
            ws.append(self._colunas)
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="073B63")
                cell.font = Font(color="FFFFFF", bold=True, name="Segoe UI",
                                 size=10)
                cell.alignment = Alignment(horizontal="center")
            for linha in linhas:
                ws.append([str(v) for v in linha])
            for i, w in enumerate(self._larguras, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter]\
                    .width = max(10, w / 7)
            wb.save(caminho)
            messagebox.showinfo("Exportado", f"Tabela salva em:\n{caminho}")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Erro ao exportar", str(exc))


# ────────────────────────────────────────────────────────────────────── #
# Estado vazio ("comece aqui") — evita cards mortos mostrando "—"        #
# ────────────────────────────────────────────────────────────────────── #
class EmptyState(Card):
    """
    Área exibida ANTES do primeiro processamento: ícone + título + dica.
    As métricas e tabelas do módulo só aparecem quando existem de fato.
    """

    def __init__(self, master, icone: str, titulo: str, dica: str, **kw):
        super().__init__(master, **kw)
        ctk.CTkLabel(self, text=icone, font=(T.FONTE, 36),
                     text_color=T.INK_3).pack(pady=(34, 2))
        ctk.CTkLabel(self, text=titulo, font=(T.FONTE, 15, "bold"),
                     text_color=T.INK).pack()
        ctk.CTkLabel(self, text=dica, font=T.F_PEQUENA, text_color=T.INK_3,
                     wraplength=560, justify="center").pack(pady=(4, 36))


# ────────────────────────────────────────────────────────────────────── #
# Toast (aviso flutuante, como no DDA)                                   #
# ────────────────────────────────────────────────────────────────────── #
class Toast:
    def __init__(self, janela: tk.Misc) -> None:
        self._janela = janela
        self._atual: Optional[ctk.CTkFrame] = None

    def show(self, mensagem: str, tipo: str = "ok", ms: int = 4000) -> None:
        self._fechar()
        cor, bg = T.STATUS_CORES.get(tipo, T.STATUS_CORES["info"])
        frame = ctk.CTkFrame(self._janela, fg_color=bg, corner_radius=10,
                             border_width=1, border_color=cor)
        ctk.CTkLabel(frame, text=mensagem, font=(T.FONTE, 12, "bold"),
                     text_color=cor).pack(padx=18, pady=10)
        frame.place(relx=0.5, rely=0.94, anchor="s")
        self._atual = frame
        self._janela.after(ms, self._fechar)

    def _fechar(self) -> None:
        if self._atual is not None:
            try:
                self._atual.destroy()
            except tk.TclError:
                pass
            self._atual = None


# ────────────────────────────────────────────────────────────────────── #
# Cabeçalho de página + botões padrão                                    #
# ────────────────────────────────────────────────────────────────────── #
def cabecalho(master, titulo: str, subtitulo: str,
              crumb: Optional[str] = None,
              icone: Optional[str] = None) -> ctk.CTkFrame:
    """
    Cabeçalho de página como FAIXA AZUL (igual à do Remessa × Banco):
    ícone opcional + breadcrumb pequeno + título branco + subtítulo claro,
    sobre o azul-assinatura #073B63. Unifica todas as telas.

    O label do título fica em `band._titulo_label` (o dashboard usa para
    atualizar a saudação conforme o horário).
    """
    band = ctk.CTkFrame(master, fg_color=T.AZUL_ESCURO, corner_radius=T.RAIO)
    inner = ctk.CTkFrame(band, fg_color="transparent")
    inner.pack(fill="x", padx=20, pady=16)

    if icone:
        ctk.CTkLabel(inner, text=icone, font=(T.FONTE, 30),
                     text_color=T.BRANCO).pack(side="left", padx=(0, 16))

    txt = ctk.CTkFrame(inner, fg_color="transparent")
    txt.pack(side="left", fill="x", expand=True)
    if crumb:
        ctk.CTkLabel(txt, text=crumb.upper(), font=T.F_CRUMB,
                     text_color=T.SIDEBAR_ITEM, anchor="w").pack(fill="x")
    lbl = ctk.CTkLabel(txt, text=titulo, font=T.F_TITULO,
                       text_color=T.BRANCO, anchor="w")
    lbl.pack(fill="x")
    ctk.CTkLabel(txt, text=subtitulo, font=T.F_SUBTITULO,
                 text_color=T.SIDEBAR_ITEM, anchor="w",
                 justify="left").pack(fill="x")
    band._titulo_label = lbl
    return band


def botao(master, texto: str, comando, primario: bool = True, **kw) -> ctk.CTkButton:
    if primario:
        kw.setdefault("fg_color", T.AZUL)
        kw.setdefault("hover_color", T.AZUL_HOVER)
        kw.setdefault("text_color", T.BRANCO)
    else:
        kw.setdefault("fg_color", T.CARD)
        kw.setdefault("hover_color", T.INFO_BG)
        kw.setdefault("text_color", T.AZUL)
        kw.setdefault("border_width", 1)
        kw.setdefault("border_color", T.AZUL)
    kw.setdefault("font", T.F_BOTAO)
    kw.setdefault("corner_radius", T.RAIO_BTN)
    kw.setdefault("height", 38)
    return ctk.CTkButton(master, text=texto, command=comando, **kw)
