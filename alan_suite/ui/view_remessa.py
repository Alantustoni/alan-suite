"""
view_remessa.py — Remessa × Banco (v1.4): UI ORIGINAL embutida na suíte.

Nada de janela separada: esta view reconstrói a tela do REMESSA × BANCO
original DENTRO da central, usando os PRÓPRIOS componentes vendorizados
do Sapiens V3 (FileDropZone, MetricsDashboard com os 7 cards,
DayStatusGrid, AttentionTable com busca/ordenação, StatusBadge,
PremiumButton) e o tema original (core/theme.py do Sapiens). O fluxo,
os textos e o comportamento são os do main_window.py original — apenas
a moldura (topbar/janela) sai, porque agora ele vive dentro da central.
"""
from __future__ import annotations

import queue
import threading
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk

from ..core import historico
from ..vendor import bootstrap

bootstrap()  # habilita imports absolutos do pacote original (core/, services/, ui/)

from core.theme import (APP_NAME, APP_TAGLINE, COLORS, FONT,  # noqa: E402
                        RADIUS)
from core.logger import log, log_exception, log_path            # noqa: E402
from core.session import Session                                 # noqa: E402
from services import (ExcelReadError, ResultadoConciliacao,      # noqa: E402
                      exportar_relatorio, parse_banco,
                      parse_sistema, reconciliar)
from ui.widgets import (FileDropZone, PremiumButton,             # noqa: E402
                        SectionTitle, StatusBadge)
from ui.dashboard import DayStatusGrid, MetricsDashboard         # noqa: E402
from ui.attention_table import AttentionTable                    # noqa: E402


class RemessaView(ctk.CTkFrame):
    """A tela original do REMESSA × BANCO, dentro da central."""

    def __init__(self, master, app) -> None:
        super().__init__(master, fg_color=COLORS["bg_main"], corner_radius=14)
        self.app = app
        self.session = Session()
        self.resultado: ResultadoConciliacao | None = None
        self._banco_files: list[str] = []
        self._sistema_files: list[str] = []
        self._queue: queue.Queue = queue.Queue()

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        self._build_header()
        self._build_main_content()
        self._build_statusbar()
        self._poll_queue()

    # ── Header azul original ─────────────────────────────────────────
    def _build_header(self):
        hdr = ctk.CTkFrame(self, fg_color=COLORS["blue_main"], corner_radius=0)
        hdr.grid(row=0, column=0, sticky="ew")
        hdr.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(hdr, text="🏦", font=FONT["5xl"]).grid(
            row=0, column=0, padx=(28, 14), pady=20)

        tf = ctk.CTkFrame(hdr, fg_color="transparent")
        tf.grid(row=0, column=1, sticky="w", pady=20)
        ctk.CTkLabel(tf, text=APP_NAME, font=FONT["3xl_b"],
                     text_color=COLORS["white"]).pack(anchor="w")
        ctk.CTkLabel(tf, text=APP_TAGLINE, font=FONT["base"],
                     text_color=COLORS["blue_light"]).pack(anchor="w",
                                                           pady=(2, 0))

        sf = ctk.CTkFrame(hdr, fg_color="transparent")
        sf.grid(row=0, column=2, padx=28, pady=20, sticky="e")
        self._badge = StatusBadge(sf, text="  AGUARDANDO  ", state="neutral")
        self._badge.pack()

    # ── Conteúdo (idêntico ao main_window original) ──────────────────
    def _build_main_content(self):
        self._scroll = ctk.CTkScrollableFrame(
            self, fg_color=COLORS["bg_main"],
            scrollbar_button_color=COLORS["border"],
            scrollbar_button_hover_color=COLORS["blue_sec"],
        )
        self._scroll.grid(row=1, column=0, sticky="nsew")
        self._scroll.grid_columnconfigure(0, weight=1)

        r = 0
        r = self._build_import_section(r)
        r = self._build_dashboard_section(r)
        r = self._build_days_section(r)
        r = self._build_table_section(r)

    def _build_import_section(self, r: int) -> int:
        SectionTitle(self._scroll, "📥  Importar Arquivos").grid(
            row=r, column=0, padx=20, pady=(20, 6), sticky="ew")
        r += 1

        fr = ctk.CTkFrame(self._scroll, fg_color="transparent")
        fr.grid(row=r, column=0, padx=20, pady=(4, 0), sticky="ew")
        fr.grid_columnconfigure(0, weight=1)
        fr.grid_columnconfigure(1, weight=1)
        r += 1

        initial = self.session.last_dir
        self._zone_banco = FileDropZone(
            fr, "🏛  Excel do BANCO", "1 arquivo (.xls / .xlsx)",
            on_select=self._on_banco_selected, multi=False,
            initial_dir=initial)
        self._zone_banco.grid(row=0, column=0, padx=(0, 8), sticky="ew")

        self._zone_sistema = FileDropZone(
            fr, "💼  Excel do SISTEMA",
            "1 a 5 arquivos — cada um representa um dia da semana",
            on_select=self._on_sistema_selected, multi=True,
            initial_dir=initial)
        self._zone_sistema.grid(row=0, column=1, padx=(8, 0), sticky="ew")

        bf = ctk.CTkFrame(self._scroll, fg_color="transparent")
        bf.grid(row=r, column=0, padx=20, pady=16, sticky="ew")
        bf.grid_columnconfigure(0, weight=1)
        r += 1

        self._btn_run = PremiumButton(
            bf, text="▶  Iniciar Conciliação",
            command=self._start_reconciliation, width=240)
        self._btn_run.grid(row=0, column=1, padx=(8, 0))

        PremiumButton(
            bf, text="↺  Limpar", command=self._clear,
            fg_color="#64748B", hover_color="#475569", width=120,
        ).grid(row=0, column=2, padx=(8, 0))
        return r

    def _build_dashboard_section(self, r: int) -> int:
        SectionTitle(self._scroll, "📊  Visão Geral").grid(
            row=r, column=0, padx=20, pady=(20, 6), sticky="ew")
        r += 1
        self._dashboard = MetricsDashboard(self._scroll)
        self._dashboard.grid(row=r, column=0, padx=20, pady=(4, 0), sticky="ew")
        return r + 1

    def _build_days_section(self, r: int) -> int:
        SectionTitle(self._scroll, "📆  Status por Dia").grid(
            row=r, column=0, padx=20, pady=(20, 6), sticky="ew")
        r += 1
        self._days_grid = DayStatusGrid(self._scroll)
        self._days_grid.grid(row=r, column=0, padx=20, pady=(4, 0), sticky="ew")
        return r + 1

    def _build_table_section(self, r: int) -> int:
        sf = ctk.CTkFrame(self._scroll, fg_color="transparent")
        sf.grid(row=r, column=0, padx=20, pady=(20, 4), sticky="ew")
        sf.grid_columnconfigure(0, weight=1)
        r += 1

        ctk.CTkLabel(sf, text="⚠️  Pontos de Atenção",
                     font=FONT["lg_b"],
                     text_color=COLORS["blue_main"]).grid(
            row=0, column=0, sticky="w")

        self._btn_export = PremiumButton(
            sf, text="📥  Exportar Excel", command=self._export,
            fg_color=COLORS["green_ok"], hover_color=COLORS["green_dark"],
            width=180)
        self._btn_export.grid(row=0, column=1)
        self._btn_export.configure(state="disabled")

        self._table = AttentionTable(self._scroll)
        self._table.grid(row=r, column=0, padx=20, pady=(4, 24), sticky="nsew")
        return r + 1

    def _build_statusbar(self):
        sb = ctk.CTkFrame(self, fg_color=COLORS["bg_topbar"], height=30,
                          corner_radius=0)
        sb.grid(row=2, column=0, sticky="ew")
        sb.grid_propagate(False)
        sb.grid_columnconfigure(0, weight=1)

        self._status_lbl = ctk.CTkLabel(
            sb, text="  Pronto. Selecione os arquivos.",
            font=FONT["xs"], text_color=COLORS["text_sec"])
        self._status_lbl.grid(row=0, column=0, padx=12, sticky="w")

        self._progress = ctk.CTkProgressBar(
            sb, width=140, height=6, fg_color=COLORS["border"],
            progress_color=COLORS["blue_sec"], corner_radius=RADIUS["sm"])
        self._progress.grid(row=0, column=1, padx=12)
        self._progress.set(0)

    # ── Callbacks (originais) ────────────────────────────────────────
    def _on_banco_selected(self, files: list[str]):
        self._banco_files = files
        if files:
            self.session.last_dir = str(Path(files[0]).parent)
        self.session.set("last_banco", files)

    def _on_sistema_selected(self, files: list[str]):
        self._sistema_files = files
        if files:
            self.session.last_dir = str(Path(files[0]).parent)
        self.session.set("last_sistema", files)

    def _set_status(self, text: str):
        self._status_lbl.configure(text=f"  {text}")

    def _clear(self):
        self._banco_files = []
        self._sistema_files = []
        self._zone_banco.clear()
        self._zone_sistema.clear()
        self._dashboard.reset()
        self._days_grid.reset()
        self._table.reset()
        self.resultado = None
        self._btn_export.configure(state="disabled")
        self._badge.update_state("  AGUARDANDO  ", "neutral")
        self._set_status("Tudo limpo. Selecione os arquivos.")
        self._progress.set(0)

    # ── Motor (original, com threading) ──────────────────────────────
    def _start_reconciliation(self):
        if not self._banco_files:
            messagebox.showwarning("Atenção", "Selecione o arquivo do Banco.")
            return
        if not self._sistema_files:
            messagebox.showwarning(
                "Atenção", "Selecione ao menos um arquivo do Sistema.")
            return

        self._btn_run.configure(state="disabled", text="⏳  Processando...")
        self._btn_export.configure(state="disabled")
        self._progress.set(0)
        self._badge.update_state("  PROCESSANDO...  ", "processing")
        self._set_status("Iniciando conciliação...")
        self._dropzones_ativas(False)
        threading.Thread(target=self._reconciliation_thread,
                         daemon=True).start()

    def _reconciliation_thread(self):
        try:
            self._queue.put(("progress", 0.1))
            self._queue.put(("status", "Lendo Excel do Banco..."))
            banco_records = parse_banco(self._banco_files[0])

            self._queue.put(("progress", 0.35))
            self._queue.put(("status", "Lendo Excel(s) do Sistema..."))
            sistema_records = []
            for i, fp in enumerate(self._sistema_files):
                sistema_records.extend(parse_sistema(fp))
                pct = 0.35 + 0.35 * (i + 1) / len(self._sistema_files)
                self._queue.put(("progress", pct))

            self._queue.put(("status", "Conciliando registros (3 passadas)..."))
            self._queue.put(("progress", 0.85))
            resultado = reconciliar(sistema_records, banco_records)

            self._queue.put(("progress", 1.0))
            self._queue.put(("resultado", resultado))
        except ExcelReadError as e:
            self._queue.put(("error_friendly", str(e)))
        except Exception as e:  # noqa: BLE001
            log_exception("Erro na conciliação", e)
            self._queue.put(("error", f"{e}\n\nVer log em:\n{log_path()}"))

    def _poll_queue(self):
        try:
            while True:
                kind, data = self._queue.get_nowait()
                if kind == "progress":
                    self._progress.set(data)
                elif kind == "status":
                    self._set_status(data)
                elif kind == "resultado":
                    self._apply_result(data)
                elif kind == "error_friendly":
                    self._on_error(data, friendly=True)
                elif kind == "error":
                    self._on_error(data, friendly=False)
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    def _dropzones_ativas(self, ativa: bool):
        for zone in (self._zone_banco, self._zone_sistema):
            for w in (zone._drop, zone._icon, zone._hint):
                try:
                    if ativa:
                        w.bind("<Button-1>",
                               lambda _e, z=zone: z._open_dialog())
                        w.bind("<Enter>", zone._hover_on)
                        w.bind("<Leave>", zone._hover_off)
                    else:
                        w.unbind("<Button-1>")
                except Exception:
                    pass

    def _on_error(self, msg: str, friendly: bool):
        self._btn_run.configure(state="normal", text="▶  Iniciar Conciliação")
        self._badge.update_state("  AGUARDANDO  ", "neutral")
        self._dropzones_ativas(True)
        title = "Arquivo inválido" if friendly else "Erro no processamento"
        messagebox.showerror(title, msg)
        self._set_status("Erro — verifique os arquivos.")

    def _apply_result(self, r: ResultadoConciliacao):
        self.resultado = r
        self._btn_run.configure(state="normal", text="▶  Iniciar Conciliação")
        self._dropzones_ativas(True)

        self._dashboard.render(r)
        self._days_grid.render(r.status_por_dia)
        self._table.render(r)

        if r.tudo_ok:
            self._badge.update_state("  ✅  TUDO OK  ", "ok")
        elif r.qtd_inconsistencias > 0 and r.qtd_divergencias > 0:
            self._badge.update_state(
                f"  ⚠️  {r.qtd_divergencias} DIV · "
                f"{r.qtd_inconsistencias} INC  ", "warning")
        elif r.qtd_inconsistencias > 0:
            self._badge.update_state(
                f"  ❌  {r.qtd_inconsistencias} INCONSISTÊNCIA(S)  ", "error")
        else:
            self._badge.update_state(
                f"  ⚠️  {r.qtd_divergencias} DIVERGÊNCIA(S)  ", "warning")

        self._btn_export.configure(state="normal")
        self._set_status(
            f"Concluído — {r.qtd_conciliados} conciliados, "
            f"{r.qtd_divergencias} divergências, "
            f"{r.qtd_inconsistencias} inconsistências.")
        log.info("Conciliação OK — conc=%d div=%d inc=%d",
                 r.qtd_conciliados, r.qtd_divergencias, r.qtd_inconsistencias)

        # integração com a central: histórico + dashboard
        historico.registrar(
            "remessa", "Conciliação",
            resumo={
                "conciliados": r.qtd_conciliados,
                "divergencias": r.qtd_divergencias,
                "inconsistencias": r.qtd_inconsistencias,
                "remessas_ausentes": len(r.remessas_ausentes),
            },
            arquivos=self._sistema_files + self._banco_files)
        self.app.atualizar_dashboard()

    # ── Exportação (original + aba "Banco sem par") ──────────────────
    def _export(self):
        if not self.resultado:
            messagebox.showinfo("Exportar", "Execute a conciliação primeiro.")
            return
        fp = filedialog.asksaveasfilename(
            title="Salvar relatório", defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialdir=self.session.last_dir,
            initialfile=f"conciliacao_{datetime.now():%Y%m%d_%H%M}.xlsx")
        if not fp:
            return
        try:
            exportar_relatorio(self.resultado, fp)
            _acrescentar_aba_banco_sem_par(self.resultado, fp)
            messagebox.showinfo("Exportado", f"Arquivo salvo em:\n{fp}")
            self._set_status(f"Exportado: {Path(fp).name}")
            historico.registrar("remessa", "Exportação de relatório",
                                saida=fp)
            self.app.atualizar_dashboard()
        except Exception as e:  # noqa: BLE001
            log_exception("Erro ao exportar", e)
            messagebox.showerror("Erro ao exportar", str(e))


# ---------------------------------------------------------------------- #
def _acrescentar_aba_banco_sem_par(resultado, caminho: str) -> None:
    """Aba extra da suíte: agendamentos do banco sem par na remessa."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(caminho)
    ws = wb.create_sheet("Banco sem par")
    ws.append(["Data", "Favorecido", "Referência", "Valor (R$)"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="073B63")
        cell.font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=10)
        cell.alignment = Alignment(horizontal="center")
    total = 0.0
    for rec in resultado.nao_conciliados_banco:
        ws.append([rec.data, rec.fornecedor, rec.referencia, rec.valor])
        total += rec.valor
    ws.append([])
    ws.append(["", "", "TOTAL", round(total, 2)])
    for col, w in zip("ABCD", (12, 46, 18, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"
    wb.save(caminho)
