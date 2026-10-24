"""
usinas.py — Organizador Usinas (reescrito, com o bug da NF corrigido).

O que o programa antigo fazia de errado (organizador.py):

  1. `str(nf).replace(".0", "")` removia ".0" de QUALQUER posição do
     número — "190.045" virava "19045" — corrompendo NFs com ponto.
  2. Depois cortava `nf[:-2]` SEMPRE, mesmo quando a NF não tinha vindo
     com sufixo — sumia dígito verdadeiro da nota.
  3. NF vazia virava a string "nan", que cortada virava "n".
  4. Data inválida virava NaT e valor vazio virava 0 — os dois entravam
     calados na planilha final.

O que esta versão faz:

  • Remove ".0" apenas do FINAL (resíduo de float do Excel) — regex ancorada.
  • O corte de sufixo é regra de negócio CONFIGURÁVEL por empresa
    (config "usinas"), com proteção de tamanho mínimo.
  • Linha com NF/data/valor inválido NÃO entra no resultado final:
    vai para a aba "Verificar" com o motivo — nada some calado.
  • Relatório com duas abas formatadas: "Tratado" e "Verificar".
"""
from __future__ import annotations

import logging
import os
import re
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ..core.utils import parse_data_br, parse_valor_br

logger = logging.getLogger(__name__)

# Resíduo de float do Excel é sempre EXATAMENTE ".0" ("123456.0").
# Com "+" a regex comia ".000" de separador de MILHAR ("198.000" → "198").
_RE_FLOAT_RESIDUO = re.compile(r"\.0$")   # só '.0' exato no fim
_RE_NAO_DIGITO    = re.compile(r"\D")


def _sem_acento(texto: str) -> str:
    """MAIÚSCULAS e sem acento — 'Siderúrgica' e 'Siderurgica' viram iguais."""
    t = unicodedata.normalize("NFKD", (texto or "").upper())
    return t.encode("ascii", "ignore").decode("ascii")

COLUNAS_SAIDA = ["EMPRESA", "NF", "VENCIMENTO", "VALOR"]


@dataclass
class ResultadoUsinas:
    """Resultado do tratamento — linhas boas, linhas para verificar, métricas."""
    linhas_ok:        List[Dict[str, Any]] = field(default_factory=list)
    linhas_verificar: List[Dict[str, Any]] = field(default_factory=list)
    total_original:   int = 0
    valor_total:      float = 0.0
    correcoes_nf:     int = 0   # NFs que a regra antiga teria corrompido
    usou_mapa_posicional: bool = False
    totais_por_grupo: Dict[str, float] = field(default_factory=dict)
    avisos:           List[str] = field(default_factory=list)

    @property
    def total_ok(self) -> int:
        return len(self.linhas_ok)

    @property
    def total_verificar(self) -> int:
        return len(self.linhas_verificar)


# Termos que denunciam uma linha de cabeçalho (planilha COM cabeçalho)
_TERMOS_CABECALHO = ("EMPRESA", "FORNECEDOR", "NF", "NOTA", "VENCIMENTO",
                     "VENC", "VALOR")


class ProcessadorUsinas:
    def __init__(self, cfg_usinas: Dict[str, Any]) -> None:
        self.padronizacao: Dict[str, List[str]] = cfg_usinas.get(
            "padronizacao_empresa", {})
        self.sufixo_padrao: int = int(cfg_usinas.get("nf_digitos_sufixo_padrao", 2))
        self.regras_especiais: List[Dict] = cfg_usinas.get("nf_regras_especiais", [])
        self.nf_tamanho_minimo: int = int(cfg_usinas.get("nf_tamanho_minimo", 3))
        # mapa posicional (planilha sem cabeçalho, colada a partir de A1)
        mp = cfg_usinas.get(
            "mapa_posicional",
            {"EMPRESA": 3, "NF": 4, "VENCIMENTO": 7, "VALOR": 9})
        self.mapa_posicional: Dict[str, int] = {k: int(v) for k, v in mp.items()}

    # ------------------------------------------------------------------ #
    def processar(self, arquivo: Path) -> ResultadoUsinas:
        import pandas as pd  # lazy — janela abre antes do pandas carregar

        logger.info("Usinas: carregando %s", arquivo.name)
        df = pd.read_excel(arquivo, dtype=str, header=None)
        if df.empty:
            raise ValueError("A planilha está vazia.")

        resultado = ResultadoUsinas()
        self._avisos_deteccao: List[str] = []
        registros, resultado.usou_mapa_posicional = self._extrair(df)
        resultado.avisos.extend(self._avisos_deteccao)

        for linha_num, bruto in registros:
            # NaN é truthy — testa o texto normalizado de CADA campo,
            # senão linha em branco no meio da planilha ia p/ 'Verificar'
            if not any(_texto(v) for v in bruto.values()):
                continue  # linha totalmente vazia
            resultado.total_original += 1

            problemas: List[str] = []

            empresa = self._padronizar_empresa(bruto["EMPRESA"])
            if not empresa:
                problemas.append("Empresa vazia")

            nf, corrigida = self._tratar_nf(bruto["NF"], empresa)
            if nf is None:
                problemas.append("NF vazia ou inválida")
            elif corrigida:
                resultado.correcoes_nf += 1

            venc = parse_data_br(bruto["VENCIMENTO"])
            if venc is None:
                problemas.append("Vencimento inválido")

            valor = parse_valor_br(bruto["VALOR"])
            if valor is None:
                problemas.append("Valor vazio ou inválido")
            elif valor == 0:
                problemas.append("Valor zerado")

            if problemas:
                resultado.linhas_verificar.append({
                    "LINHA":      linha_num,
                    "EMPRESA":    bruto["EMPRESA"],
                    "NF":         bruto["NF"],
                    "VENCIMENTO": bruto["VENCIMENTO"],
                    "VALOR":      _texto(bruto["VALOR"]),
                    "MOTIVO":     " | ".join(problemas),
                })
            else:
                resultado.linhas_ok.append({
                    "EMPRESA":    empresa,
                    "NF":         nf,
                    "VENCIMENTO": venc,
                    "VALOR":      valor,
                })
                resultado.valor_total += valor  # type: ignore[operator]

        resultado.linhas_ok.sort(
            key=lambda r: (r["VENCIMENTO"], r["EMPRESA"],
                           (len(r["NF"]), r["NF"])))  # NF numérica, não lexicográfica

        # totais por grupo de usina (USIMINAS/ARMAR/CSN/SOLUÇÕES/…)
        totais: Dict[str, float] = {}
        for l in resultado.linhas_ok:
            totais[l["EMPRESA"]] = totais.get(l["EMPRESA"], 0.0) + l["VALOR"]
        resultado.totais_por_grupo = totais

        logger.info(
            "Usinas: %d ok, %d para verificar (%d NFs protegidas da regra antiga)",
            resultado.total_ok, resultado.total_verificar, resultado.correcoes_nf,
        )
        return resultado

    # ------------------------------------------------------------------ #
    # Extração de colunas: cabeçalho OU mapa posicional (colar em A1)      #
    # ------------------------------------------------------------------ #
    def _extrair(self, df) -> Tuple[List[Tuple[int, Dict[str, Any]]], bool]:
        """
        Devolve (lista de (linha_excel, {EMPRESA,NF,VENCIMENTO,VALOR}), usou_posicional).
        Detecta se a 1ª linha é cabeçalho; se não for, usa o mapa posicional
        das usinas (como o DDA, porém SEM a coluna FILIAL — a usina vem na
        coluna D). A coluna da usina é confirmada pelo conteúdo.
        """
        import pandas as pd  # noqa: F401

        primeira = " ".join(_texto(v).upper() for v in df.iloc[0].tolist())
        tem_cabecalho = sum(t in primeira for t in _TERMOS_CABECALHO) >= 2

        if tem_cabecalho:
            colunas = [_texto(c).upper() for c in df.iloc[0].tolist()]
            col_idx = self._mapear_colunas(colunas)
            faltando = [k for k, v in col_idx.items() if v is None]
            if faltando:
                raise ValueError(
                    "Colunas obrigatórias não encontradas: "
                    + ", ".join(faltando)
                    + ".\nA planilha precisa ter EMPRESA/FORNECEDOR, NF, "
                    "VENCIMENTO e VALOR — ou cole sem cabeçalho no layout do DDA.")
            corpo = df.iloc[1:]
            offset = 2   # +1 cabeçalho, +1 base do Excel
            idx = col_idx
        else:
            ncols = df.shape[1]
            corpo = df
            offset = 1
            # As colunas são descobertas pelo CONTEÚDO (robusto a variações de
            # layout — filial a menos, colunas a mais/menos). O mapa posicional
            # é só um palpite inicial (fallback).
            idx = dict(self.mapa_posicional)
            col_usina = self._detectar_coluna_usina(corpo)
            if col_usina is not None:
                idx["EMPRESA"] = col_usina
                if col_usina + 1 < ncols:      # a NF costuma vir logo após a usina
                    idx["NF"] = col_usina + 1
            col_venc = self._detectar_coluna_venc(corpo)
            if col_venc is not None:
                idx["VENCIMENTO"] = col_venc
            col_valor = self._detectar_coluna_valor(corpo, evitar={col_venc})
            if col_valor is not None:
                idx["VALOR"] = col_valor

            faltam = [n for n in ("EMPRESA", "NF", "VENCIMENTO", "VALOR")
                      if idx.get(n, 99) >= ncols]
            if faltam:
                raise ValueError(
                    f"A planilha colada tem só {ncols} coluna(s) e não consegui "
                    f"localizar: {', '.join(faltam)}.\n"
                    "Cole a partir de A1, sem cabeçalho, com as colunas do "
                    "relatório das usinas (usina, NF, vencimento e valor).")

        registros: List[Tuple[int, Dict[str, Any]]] = []
        for i, (_, row) in enumerate(corpo.iterrows()):
            registros.append((i + offset, {
                "EMPRESA":    _texto(row.iloc[idx["EMPRESA"]]),
                "NF":         _texto(row.iloc[idx["NF"]]),
                "VENCIMENTO": _texto(row.iloc[idx["VENCIMENTO"]]),
                "VALOR":      row.iloc[idx["VALOR"]],
            }))
        return registros, (not tem_cabecalho)

    def _detectar_coluna_usina(self, corpo) -> Optional[int]:
        """
        Descobre QUAL coluna contém o nome da usina, pelo conteúdo.

        Sem cabeçalho não dá para confiar numa posição fixa (foi o que
        colocava tudo em USIMINAS). Varremos cada coluna numa amostra das
        linhas e escolhemos a que gera MAIS grupos distintos reconhecidos
        (USIMINAS/ARMAR/CSN/SOLUÇÕES…). Desempate pela quantidade de
        células que caem num grupo conhecido.

        Retorna o índice da coluna, ou None se nenhuma coluna reconhecer
        usinas (aí o mapa posicional configurado é mantido).
        """
        ncols = corpo.shape[1]
        amostra = corpo.head(300)
        melhor: Optional[int] = None
        melhor_score = (0, 0)   # (grupos_distintos, celulas_reconhecidas)
        for c in range(ncols):
            grupos = set()
            reconhecidas = 0
            for v in amostra.iloc[:, c].tolist():
                g = self._padronizar_empresa(_texto(v))
                if g in self.padronizacao:      # caiu num grupo conhecido
                    grupos.add(g)
                    reconhecidas += 1
            score = (len(grupos), reconhecidas)
            if score > melhor_score:
                melhor_score, melhor = score, c
        return melhor if melhor_score[1] > 0 else None

    def _detectar_coluna_venc(self, corpo) -> Optional[int]:
        """
        Acha a coluna de VENCIMENTO pelo conteúdo (a que tem datas).

        Pode haver mais de uma coluna de data (vencimento, data-limite no
        futuro, data de emissão no passado). Regra:
          1. se a coluna do MAPA POSICIONAL configurado é candidata, ela
             vence (evita trocar vencimento por emissão em silêncio);
          2. senão, mediana mais CEDO (descarta a data-limite futura).
        Com 2+ candidatas registra um aviso — nada de troca silenciosa.
        """
        ncols = corpo.shape[1]
        amostra = corpo.head(300)
        minimo = max(3, int(0.4 * len(amostra)))
        candidatos = []   # (mediana, quantidade, coluna)
        for c in range(ncols):
            datas = [d for d in (parse_data_br(_texto(v))
                                 for v in amostra.iloc[:, c].tolist())
                     if d is not None]
            if len(datas) >= minimo:
                datas.sort()
                candidatos.append((datas[len(datas) // 2], len(datas), c))
        if not candidatos:
            return None
        colunas_cand = [c for _, _, c in candidatos]
        preferida = self.mapa_posicional.get("VENCIMENTO")
        if preferida in colunas_cand:
            escolhida = preferida
        else:
            candidatos.sort(key=lambda t: (t[0], -t[1]))   # mediana mais cedo
            escolhida = candidatos[0][2]
        if len(colunas_cand) > 1:
            self._avisar(
                f"Mais de uma coluna de datas na planilha (colunas "
                f"{colunas_cand}) — usei a coluna {escolhida} como "
                "vencimento. Confira as datas no resultado.")
        return escolhida

    def _detectar_coluna_valor(self, corpo, evitar=None) -> Optional[int]:
        """
        Acha a coluna de VALOR: números com CENTAVOS (2 casas após , ou .).
        Exigir 2 casas separa o valor (126.908,69) de NF/código com ponto
        de milhar ("190.045"). Se a coluna do mapa posicional qualificar,
        ela vence; senão, a candidata mais à direita.
        """
        import re as _re
        evitar = {e for e in (evitar or set()) if e is not None}
        ncols = corpo.shape[1]
        amostra = corpo.head(300)
        contagens: Dict[int, int] = {}
        for c in range(ncols):
            if c in evitar:
                continue
            com_centavos = 0
            for v in amostra.iloc[:, c].tolist():
                t = _texto(v)
                if not t or not _re.search(r"[.,]\d{2}$", t):
                    continue                     # exige 2 casas decimais
                val = parse_valor_br(t)
                if val is not None and val > 0:
                    com_centavos += 1
            if com_centavos > 0:
                contagens[c] = com_centavos
        if not contagens:
            return None
        preferida = self.mapa_posicional.get("VALOR")
        maximo = max(contagens.values())
        # a coluna configurada vence quando está entre as boas candidatas
        if preferida in contagens and contagens[preferida] >= maximo * 0.9:
            return preferida
        # senão: a mais à direita entre as de maior contagem
        return max(c for c, n in contagens.items() if n == maximo)

    def _avisar(self, msg: str) -> None:
        avisos = getattr(self, "_avisos_deteccao", None)
        if avisos is not None and msg not in avisos:
            avisos.append(msg)
        logger.warning("Usinas: %s", msg)

    def _mapear_colunas(self, colunas) -> Dict[str, Optional[int]]:
        """
        Índice das colunas pelo NOME (planilha com cabeçalho).

        A prioridade é do ALIAS, não da posição: procura-se cada alias em
        ordem (match exato primeiro, depois 'contém'). Sem isso, uma coluna
        "DATA EMISSÃO" à esquerda de "VENCIMENTO" capturava o vencimento.
        """
        aliases = {
            "EMPRESA":    ["FORNECEDOR", "EMPRESA", "CLIENTE"],
            "NF":         ["NF", "NOTA", "DOCUMENTO"],
            "VENCIMENTO": ["VENCIMENTO", "VENC", "DATA"],
            "VALOR":      ["VALOR", "VLR", "TOTAL"],
        }
        resultado: Dict[str, Optional[int]] = {}
        for alvo, opcoes in aliases.items():
            achada = None
            for op in opcoes:                    # prioridade do alias
                for i, col in enumerate(colunas):
                    if col == op:                # 1º passe: match exato
                        achada = i
                        break
                if achada is not None:
                    break
                for i, col in enumerate(colunas):
                    if op in col:                # 2º passe: contém
                        achada = i
                        break
                if achada is not None:
                    break
            resultado[alvo] = achada
        return resultado

    def _padronizar_empresa(self, texto: str) -> str:
        """
        Classifica o nome da usina num grupo, ignorando acentos.

        A ordem do dicionário de padronização importa: o 1º grupo cujo
        termo aparece no nome vence. Por isso SOLUÇÕES vem antes de
        USIMINAS na config — 'SOLUÇÕES USIMINAS' deve virar SOLUÇÕES, não
        ser engolido pelo termo genérico 'USINAS' de USIMINAS.
        """
        t = _sem_acento(texto).strip()
        for padrao, termos in self.padronizacao.items():
            if any(_sem_acento(termo) in t for termo in termos):
                return padrao
        return t

    # ------------------------------------------------------------------ #
    # A CORREÇÃO DO BUG                                                    #
    # ------------------------------------------------------------------ #
    def _tratar_nf(self, nf_bruta: str, empresa: str) -> Tuple[Optional[str], bool]:
        """
        Retorna (nf_tratada, regra_antiga_teria_corrompido).

        Passos:
          1. Descarta vazio/nan — vai para "Verificar" (antes virava "n").
          2. Remove APENAS o resíduo de float ".0" no FIM da string
             (antes o replace pegava ".0" em qualquer posição).
          3. Remove outros separadores não numéricos (ponto de milhar).
          4. Aplica o corte de sufixo configurado por empresa, mas só
             quando sobra pelo menos `nf_tamanho_minimo` dígitos.
        """
        texto = (nf_bruta or "").strip()
        if not texto or texto.lower() in ("nan", "none", "nat"):
            return None, False

        # A regra antiga corromperia esta NF? (para métrica de auditoria)
        antiga_corrompia = (".0" in texto[:-2]) if len(texto) > 2 else False

        # 1) resíduo de float SÓ no final: "123456.0" → "123456"
        texto = _RE_FLOAT_RESIDUO.sub("", texto)
        # 2) qualquer outro não-dígito (ponto de milhar, espaço, hífen)
        digitos = _RE_NAO_DIGITO.sub("", texto)
        if not digitos:
            return None, False

        # 3) corte de sufixo (regra de negócio, explícita e configurável).
        #    Cada regra pode casar por empresa (opcional; vazio/"*" = todas),
        #    por INÍCIO (nf_comeca_com) e/ou por FIM (nf_termina_com) da NF.
        corte = self.sufixo_padrao
        for regra in self.regras_especiais:
            emp_regra = str(regra.get("empresa", "")).upper()
            if emp_regra and emp_regra not in ("*", empresa):
                continue
            comeca = str(regra.get("nf_comeca_com", ""))
            termina = str(regra.get("nf_termina_com", ""))
            if comeca and not digitos.startswith(comeca):
                continue
            if termina and not digitos.endswith(termina):
                continue
            corte = int(regra.get("digitos_sufixo", corte))
            break

        if corte > 0 and len(digitos) - corte >= self.nf_tamanho_minimo:
            digitos = digitos[:-corte]

        return digitos, antiga_corrompia


# ---------------------------------------------------------------------- #
# Exportação                                                              #
# ---------------------------------------------------------------------- #
def exportar(resultado: ResultadoUsinas, arquivo_saida: Path) -> None:
    """
    Gera o Excel com abas 'Tratado' e 'Verificar', formatadas.
    Escrita ATÔMICA: monta em .tmp e faz os.replace no fim — se algo
    falhar no meio, o arquivo final não fica corrompido.
    """
    import pandas as pd
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    azul   = "073B63"
    zebra  = "F2F8FE"
    borda  = Side(style="thin", color="B8D4F3")
    box    = Border(left=borda, right=borda, top=borda, bottom=borda)

    df_ok = pd.DataFrame(resultado.linhas_ok, columns=COLUNAS_SAIDA)
    df_vf = pd.DataFrame(
        resultado.linhas_verificar,
        columns=["LINHA", "EMPRESA", "NF", "VENCIMENTO", "VALOR", "MOTIVO"],
    )

    saida = Path(arquivo_saida)
    tmp_fd, tmp_path = tempfile.mkstemp(dir=saida.parent, suffix=".tmp.xlsx")
    os.close(tmp_fd)
    try:
        _escrever_xlsx(tmp_path, df_ok, df_vf, azul, zebra, box,
                       Alignment, Font, PatternFill)
        os.replace(tmp_path, saida)
    except Exception:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise
    logger.info("Usinas: exportado %s", saida.name)


def _escrever_xlsx(caminho, df_ok, df_vf, azul, zebra, box,
                   Alignment, Font, PatternFill) -> None:
    import pandas as pd
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        df_ok.to_excel(writer, index=False, sheet_name="Tratado")
        df_vf.to_excel(writer, index=False, sheet_name="Verificar")

        for nome, df in (("Tratado", df_ok), ("Verificar", df_vf)):
            ws = writer.book[nome]
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor=azul)
                cell.font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=10)
                cell.alignment = Alignment(horizontal="center")
                cell.border = box
            for r in range(2, ws.max_row + 1):
                for cell in ws[r]:
                    cell.border = box
                    if r % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor=zebra)
            larguras = {"Tratado": [24, 16, 14, 16],
                        "Verificar": [8, 24, 16, 14, 16, 42]}[nome]
            for i, w in enumerate(larguras, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        ws = writer.book["Tratado"]
        # NF como NÚMERO de verdade — remove o alerta "número armazenado como
        # texto" (o triângulo verde). NF com zero à esquerda continua texto,
        # para não perder o zero.
        for cell in ws["B"][1:]:
            txt = str(cell.value).strip() if cell.value is not None else ""
            if txt.isdigit() and (len(txt) == 1 or not txt.startswith("0")):
                cell.value = int(txt)
                cell.number_format = "0"
            else:
                cell.number_format = "@"
        for cell in ws["C"][1:]:
            cell.number_format = "DD/MM/YYYY"
        for cell in ws["D"][1:]:
            cell.number_format = "#,##0.00"


def _texto(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none", "nat") else s
